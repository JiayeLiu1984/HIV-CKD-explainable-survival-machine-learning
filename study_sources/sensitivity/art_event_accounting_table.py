from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


REFERENCE = Path(r"__CKD_LOCAL_INPUT__/多中心艾滋ML/修稿数据/Table_S1_CKD敏感性结局QC_精简版.docx")
OUTPUT = Path(r"__CKD_LOCAL_INPUT__/Documents/深度学习/outputs/comment1-art-current-exposure-20260904/Table_S1_ART相关CKD误分类敏感性分析.docx")


def set_font(run, size=11, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def set_cell_shading(cell, fill=None):
    tcpr = cell._tc.get_or_add_tcPr()
    old = tcpr.find(qn("w:shd"))
    if old is not None:
        tcpr.remove(old)
    if fill:
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), fill)
        tcpr.append(shd)


def set_cell_borders(cell, top=None, bottom=None):
    tcpr = cell._tc.get_or_add_tcPr()
    existing = tcpr.find(qn("w:tcBorders"))
    if existing is not None:
        tcpr.remove(existing)
    borders = OxmlElement("w:tcBorders")
    for edge, spec in (("top", top), ("bottom", bottom)):
        if spec:
            el = OxmlElement(f"w:{edge}")
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), str(spec))
            el.set(qn("w:space"), "0")
            el.set(qn("w:color"), "000000")
            borders.append(el)
    if len(borders):
        tcpr.append(borders)


def set_cell_margins(cell, top=35, start=80, bottom=35, end=80):
    tcpr = cell._tc.get_or_add_tcPr()
    tc_mar = tcpr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tcpr.append(tc_mar)
    for m, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def keep_row(row):
    trpr = row._tr.get_or_add_trPr()
    if trpr.find(qn("w:cantSplit")) is None:
        trpr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    trpr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trpr.append(header)


def write_cell(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT, indent=0, size=11):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.left_indent = Inches(indent)
    r = p.add_run(text)
    set_font(r, size=size, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


doc = Document(REFERENCE)
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(1.0)
section.bottom_margin = Inches(1.0)
section.left_margin = Inches(1.0)
section.right_margin = Inches(1.0)

caption = doc.paragraphs[0]
caption.clear()
caption.alignment = WD_ALIGN_PARAGRAPH.LEFT
caption.paragraph_format.space_before = Pt(0)
caption.paragraph_format.space_after = Pt(6)
caption.paragraph_format.line_spacing = 1.0
run = caption.add_run("Table S1. Sensitivity analysis of ART-related CKD outcome misclassification")
set_font(run, size=12, bold=True)

table = doc.tables[0]
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
table.columns[0].width = Inches(4.75)
table.columns[1].width = Inches(1.75)

# Keep the source header row, remove all old result rows, then extend it.
for row in list(table.rows)[1:]:
    table._tbl.remove(row._tr)

rows = [
    ("section", "CKD outcome in the overall cohort", ""),
    ("data", "Original CKD event", "2,491 (7.81)"),
    ("data", "Clinical-diagnosis CKD", "1,796 (5.63)"),
    ("data", "eGFR-based CKD", "695 (2.18)"),
    ("section", "Relevant ART exposure among eGFR-based CKD events (n = 695)", ""),
    ("data", "INSTI", "206 (29.64)"),
    ("data", "Boosted PI", "262 (37.70)"),
    ("data", "Rilpivirine", "2 (0.29)"),
    ("data", "Any relevant ART", "469 (67.48)"),
    ("section", "Sensitivity outcome", ""),
    ("data", "Sensitivity CKD event", "2,022 (6.34)"),
    ("data", "Original CKD censored at the event time", "469 (18.83)"),
    ("section", "Time from relevant ART initiation or switch to CKD (n = 469)", ""),
    ("data", "≤3 months", "92 (19.62)"),
    ("data", ">3–6 months", "73 (15.57)"),
    ("data", ">6 months", "304 (64.82)"),
]

header = table.rows[0]
write_cell(header.cells[0], "Outcome measure", bold=True, size=11)
write_cell(header.cells[1], "Overall cohort\n(N = 31,911), n (%)", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
set_cell_borders(header.cells[0], top=12, bottom=12)
set_cell_borders(header.cells[1], top=12, bottom=12)
repeat_header(header)
keep_row(header)

for kind, label, value in rows:
    row = table.add_row()
    keep_row(row)
    for i, width in enumerate((Inches(4.75), Inches(1.75))):
        row.cells[i].width = width
    if kind == "section":
        write_cell(row.cells[0], label, bold=True, size=11)
        write_cell(row.cells[1], "", bold=True, size=11)
        set_cell_shading(row.cells[0], "F2F2F2")
        set_cell_shading(row.cells[1], "F2F2F2")
        set_cell_borders(row.cells[0])
        set_cell_borders(row.cells[1])
    else:
        write_cell(row.cells[0], label, indent=0.12, size=11)
        write_cell(row.cells[1], value, align=WD_ALIGN_PARAGRAPH.CENTER, size=11)
        set_cell_shading(row.cells[0])
        set_cell_shading(row.cells[1])
        set_cell_borders(row.cells[0])
        set_cell_borders(row.cells[1])

last = table.rows[-1]
set_cell_borders(last.cells[0], bottom=12)
set_cell_borders(last.cells[1], bottom=12)

note = doc.paragraphs[1]
note.clear()
note.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
note.paragraph_format.space_before = Pt(6)
note.paragraph_format.space_after = Pt(0)
note.paragraph_format.line_spacing = 1.0
note_text = (
    "Values are n (%). Percentages for CKD outcomes are based on the full cohort; percentages for ART classes are based on "
    "eGFR-based CKD events (n = 695); the censored proportion is based on original CKD events (n = 2,491); and timing "
    "percentages are based on potentially affected CKD events (n = 469). Relevant ART exposure was defined as current use "
    "of an INSTI, boosted PI, or rilpivirine at CKD onset according to regimen start and switch records; ever-use was not used. "
    "In the sensitivity analysis, these potentially affected eGFR-based CKD events were censored at the original event time. "
    "One patient used both an INSTI and a boosted PI; therefore, ART class counts are not mutually exclusive. CKD, chronic "
    "kidney disease; eGFR, estimated glomerular filtration rate; ART, antiretroviral therapy; INSTI, integrase strand transfer "
    "inhibitor; PI, protease inhibitor."
)
r = note.add_run(note_text)
set_font(r, size=9.5)

# Keep caption with the table and prevent accidental empty trailing pages.
caption.paragraph_format.keep_with_next = True
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
