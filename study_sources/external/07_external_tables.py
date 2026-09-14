# -*- coding: utf-8 -*-

# =============================================================================
# 重庆独立地理外部验证
# FINAL CORE PERFORMANCE TABLE WITH 95% CI
#
# 最终表：
#   Patients at risk
#   CKD events within 5 years
#   5-year Uno C-index (95% CI)
#   iAUC (95% CI)
#   IBS (95% CI)
#   5-year ICI (95% CI)
#
# IBS / ICI:
#   patient-level bootstrap = 1000
#   每次bootstrap内重新执行5-fold cross-fitted
#   intercept-only recalibration
#
# 因此CI同时反映：
#   - 外部患者抽样不确定性
#   - center recalibration参数不确定性
#
# C-index / iAUC:
#   继续使用Step8 primary frozen external-validation bootstrap结果
#
# 数值：
#   三位有效数字
#   CI格式：
#   0.818（0.733，0.897）
# =============================================================================


from pathlib import Path
import time
import warnings

import numpy as np
import pandas as pd

from sklearn.preprocessing import SplineTransformer
from sklearn.linear_model import LogisticRegression
from sklearn.pipeline import make_pipeline

from sksurv.metrics import brier_score
from sksurv.nonparametric import kaplan_meier_estimator
from sksurv.util import Surv

from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    Side,
)


# =============================================================================
# 1. PATHS
# =============================================================================

BASE = Path(
    "__CKD_WORKDIR__"
)

STEP1 = (
    BASE
    / "rolling_5y_step1_new_split"
)

STEP6 = (
    BASE
    / "重庆外部验证_step6_LSTM_v2"
)

STEP7R = (
    BASE
    / "重庆外部验证_step7R_FINAL_fold_specific"
)

STEP8 = (
    BASE
    / "重庆外部验证_step8_FINAL_performance_FIXED"
)

RECAL = (
    BASE
    / "重庆外部验证_secondary_crossfit_recalibration"
)

FINAL_DIR = (
    BASE
    / "重庆外部验证_FINAL_5Figures_Table"
)

FINAL_DIR.mkdir(
    parents=True,
    exist_ok=True,
)


# =============================================================================
# 2. FILES
# =============================================================================

LANDMARK_FILE = (
    STEP8
    / "04_landmark_performance_with_95CI.csv"
)

METADATA_FILE = (
    STEP6
    / "super_landmark_external_metadata.csv"
)

ORIGINAL_HAZARD_FILE = (
    STEP7R
    / "external_calibrated_hazard_long.npy"
)

FUTURE_EVENT_FILE = (
    STEP6
    / "external_future_event_long.npy"
)

FUTURE_AT_RISK_FILE = (
    STEP6
    / "external_future_at_risk_long.npy"
)

PATIENT_FOLD_FILE = (
    RECAL
    / "patient_recalibration_folds.csv"
)

RECAL_RISK_FILE = (
    RECAL
    / "external_crossfit_recalibrated_risk_long.npy"
)

BRIER_FILE = (
    FINAL_DIR
    / "SourceData_Figure3_Brier.csv"
)


# =============================================================================
# 3. PARAMETERS
# =============================================================================

PRIMARY_LANDMARK_INDEX = np.asarray(
    [0, 1, 3, 5],
    dtype=int,
)

PRIMARY_LANDMARKS = np.asarray(
    [0.0, 1.0, 3.0, 5.0],
    dtype=float,
)

COLUMN_NAMES = {
    0.0: "Baseline (0 y)",
    1.0: "ART year 1",
    3.0: "ART year 3",
    5.0: "ART year 5",
}

LANDMARK_MONTHS = np.asarray(
    [0, 12, 24, 36, 48, 60],
    dtype=float,
)

METRIC_TIMES = np.asarray(
    [
        6,
        12,
        18,
        24,
        30,
        36,
        42,
        48,
        54,
        59.999,
    ],
    dtype=float,
)

N_FUTURE = 10
N_FOLDS = 5

BOOTSTRAP_REPS = 1000
RANDOM_SEED = 20260824

EPS = 1e-7


# =============================================================================
# 4. FILE CHECK
# =============================================================================

required_files = [
    LANDMARK_FILE,
    METADATA_FILE,
    ORIGINAL_HAZARD_FILE,
    FUTURE_EVENT_FILE,
    FUTURE_AT_RISK_FILE,
    PATIENT_FOLD_FILE,
    RECAL_RISK_FILE,
    BRIER_FILE,
    STEP1 / "development_idx.npy",
    STEP1 / "event.npy",
    STEP1 / "observed_time_month.npy",
]

missing = [
    str(path)
    for path in required_files
    if not path.exists()
]

if missing:
    raise FileNotFoundError(
        "缺少必要文件：\n"
        + "\n".join(missing)
    )


# =============================================================================
# 5. READ DATA
# =============================================================================

landmark_df = pd.read_csv(
    LANDMARK_FILE,
    encoding="utf-8-sig",
)

metadata = pd.read_csv(
    METADATA_FILE,
    encoding="utf-8-sig",
)

patient_fold_df = pd.read_csv(
    PATIENT_FOLD_FILE,
    encoding="utf-8-sig",
)

brier_point_df = pd.read_csv(
    BRIER_FILE,
    encoding="utf-8-sig",
)

original_hazard = np.load(
    ORIGINAL_HAZARD_FILE
).astype(float)

future_event = np.load(
    FUTURE_EVENT_FILE
).astype(float)

future_at_risk = np.load(
    FUTURE_AT_RISK_FILE
).astype(bool)

recal_risk_point = np.load(
    RECAL_RISK_FILE
).astype(float)


# =============================================================================
# 6. BASIC QC
# =============================================================================

LONG_N = len(metadata)

expected_shape = (
    LONG_N,
    N_FUTURE,
)

for name, array in [
    ("original_hazard", original_hazard),
    ("future_event", future_event),
    ("future_at_risk", future_at_risk),
    ("recal_risk_point", recal_risk_point),
]:

    if array.shape != expected_shape:
        raise ValueError(
            f"{name}形状错误：{array.shape}"
        )


patient_index_long = (
    metadata[
        "local_patient_index"
    ]
    .to_numpy(int)
)

landmark_index_long = (
    metadata[
        "landmark_index"
    ]
    .to_numpy(int)
)

N_PATIENT = int(
    patient_index_long.max()
    + 1
)


if len(patient_fold_df) != N_PATIENT:
    raise ValueError(
        "patient_recalibration_folds患者数异常。"
    )


patient_fold = np.full(
    N_PATIENT,
    -1,
    dtype=int,
)

patient_fold[
    patient_fold_df[
        "local_patient_index"
    ].to_numpy(int)
] = (
    patient_fold_df[
        "recalibration_fold"
    ].to_numpy(int)
)

if np.any(patient_fold < 0):
    raise ValueError(
        "存在未分配recalibration fold的患者。"
    )


# =============================================================================
# 7. DEVELOPMENT IPCW REFERENCES
# =============================================================================

development_idx = np.load(
    STEP1 / "development_idx.npy"
).astype(int)

development_event_all = np.load(
    STEP1 / "event.npy"
).astype(int)

development_time_all = np.load(
    STEP1 / "observed_time_month.npy"
).astype(float)


development_event = (
    development_event_all[
        development_idx
    ]
)

development_time = (
    development_time_all[
        development_idx
    ]
)


def make_surv(
    event,
    time,
):

    return Surv.from_arrays(
        event=np.asarray(
            event,
            dtype=bool,
        ),
        time=np.asarray(
            time,
            dtype=float,
        ),
    )


development_reference = {}


for landmark_index in PRIMARY_LANDMARK_INDEX:

    landmark_month = (
        LANDMARK_MONTHS[
            landmark_index
        ]
    )

    eligible = (
        development_time
        > landmark_month
        + EPS
    )

    event_original = (
        development_event[
            eligible
        ]
    )

    residual = (
        development_time[
            eligible
        ]
        - landmark_month
    )

    event_60 = (
        (event_original == 1)
        &
        (
            residual
            <= 60.0
            + EPS
        )
    )

    reference_time = np.minimum(
        residual,
        61.0,
    )

    development_reference[
        landmark_index
    ] = make_surv(
        event_60,
        reference_time,
    )


# =============================================================================
# 8. MATHEMATICAL FUNCTIONS
# =============================================================================

def clip_probability(p):

    return np.clip(
        np.asarray(
            p,
            dtype=float,
        ),
        EPS,
        1.0 - EPS,
    )


def logit(p):

    p = clip_probability(p)

    return (
        np.log(p)
        -
        np.log1p(-p)
    )


def expit(x):

    x = np.asarray(
        x,
        dtype=float,
    )

    result = np.empty_like(x)

    positive = (
        x >= 0
    )

    result[
        positive
    ] = (
        1.0
        /
        (
            1.0
            +
            np.exp(
                -x[
                    positive
                ]
            )
        )
    )

    exp_x = np.exp(
        x[
            ~positive
        ]
    )

    result[
        ~positive
    ] = (
        exp_x
        /
        (
            1.0
            +
            exp_x
        )
    )

    return result


def hazard_to_risk(hazard):

    survival = np.cumprod(
        1.0
        - clip_probability(
            hazard
        ),
        axis=1,
    )

    return (
        1.0
        - survival
    )


# =============================================================================
# 9. WEIGHTED INTERCEPT-ONLY RECALIBRATION
#
# bootstrap中患者可能被抽中多次，
# multiplicity作为sample weight。
# =============================================================================

def fit_weighted_intercept(
    hazard,
    outcome,
    weight,
):

    hazard = clip_probability(
        hazard
    )

    outcome = np.asarray(
        outcome,
        dtype=float,
    )

    weight = np.asarray(
        weight,
        dtype=float,
    )

    keep = (
        weight > 0
    )

    hazard = hazard[
        keep
    ]

    outcome = outcome[
        keep
    ]

    weight = weight[
        keep
    ]


    if len(hazard) == 0:
        raise ValueError(
            "无有效训练数据。"
        )


    event_weight = np.sum(
        weight * outcome
    )

    total_weight = np.sum(
        weight
    )


    if event_weight <= 0:
        raise ValueError(
            "bootstrap训练样本无事件。"
        )


    if event_weight >= total_weight:
        raise ValueError(
            "bootstrap训练样本全部为事件。"
        )


    offset = logit(
        hazard
    )


    def score(delta):

        probability = expit(
            offset
            + delta
        )

        return float(
            np.sum(
                weight
                *
                (
                    probability
                    -
                    outcome
                )
            )
        )


    lower = -15.0
    upper = 15.0


    for _ in range(120):

        middle = (
            lower
            + upper
        ) / 2.0

        if score(middle) > 0:
            upper = middle

        else:
            lower = middle


    return float(
        (
            lower
            + upper
        )
        / 2.0
    )


# =============================================================================
# 10. BUILD ONE BOOTSTRAP CROSS-FITTED RECALIBRATION
# =============================================================================

def build_bootstrap_recalibrated_risk(
    patient_counts,
):

    recal_hazard = np.full_like(
        original_hazard,
        np.nan,
        dtype=float,
    )


    for fold_id in range(
        N_FOLDS
    ):

        train_patient = (
            patient_fold
            != fold_id
        )

        test_patient = (
            patient_fold
            == fold_id
        )


        for landmark_index in (
            PRIMARY_LANDMARK_INDEX
        ):

            train_rows = np.where(
                train_patient[
                    patient_index_long
                ]
                &
                (
                    landmark_index_long
                    == landmark_index
                )
            )[0]


            test_rows = np.where(
                test_patient[
                    patient_index_long
                ]
                &
                (
                    landmark_index_long
                    == landmark_index
                )
            )[0]


            # -------------------------------------------------------------
            # interval-level mask
            # -------------------------------------------------------------

            interval_mask = (
                future_at_risk[
                    train_rows
                ]
            )


            hazard_train = (
                original_hazard[
                    train_rows
                ][
                    interval_mask
                ]
            )


            outcome_train = (
                future_event[
                    train_rows
                ][
                    interval_mask
                ]
            )


            # -------------------------------------------------------------
            # 每个origin的bootstrap multiplicity，
            # 扩展到10个interval后再按at-risk mask筛选
            # -------------------------------------------------------------

            row_weights = (
                patient_counts[
                    patient_index_long[
                        train_rows
                    ]
                ]
            )


            interval_weights = (
                np.repeat(
                    row_weights[
                        :,
                        None
                    ],
                    N_FUTURE,
                    axis=1,
                )[
                    interval_mask
                ]
            )


            delta = fit_weighted_intercept(
                hazard_train,
                outcome_train,
                interval_weights,
            )


            recal_hazard[
                test_rows
            ] = expit(
                logit(
                    original_hazard[
                        test_rows
                    ]
                )
                + delta
            )


    # -------------------------------------------------------------------------
    # 只需要主Landmark；
    # 其它2/4年Landmark可以保留NaN
    # -------------------------------------------------------------------------

    return hazard_to_risk(
        np.where(
            np.isfinite(
                recal_hazard
            ),
            recal_hazard,
            0.5,
        )
    ), recal_hazard


# =============================================================================
# 11. BRIER / IBS
# =============================================================================

def calculate_ibs(
    landmark_index,
    bootstrap_rows,
    risk,
):

    event = (
        metadata.iloc[
            bootstrap_rows
        ][
            "event_within_60m"
        ]
        .to_numpy(int)
        .astype(bool)
    )


    time = (
        metadata.iloc[
            bootstrap_rows
        ][
            "analysis_time_month"
        ]
        .to_numpy(float)
    )


    risk = np.asarray(
        risk,
        dtype=float,
    )


    y_test = make_surv(
        event,
        time,
    )


    y_train = (
        development_reference[
            landmark_index
        ]
    )


    survival = (
        1.0
        -
        risk
    )


    test_min = float(
        time.min()
    )

    test_max = float(
        time.max()
    )


    brier = np.full(
        N_FUTURE,
        np.nan,
        dtype=float,
    )


    for position, metric_time in enumerate(
        METRIC_TIMES
    ):

        if metric_time < test_min:

            # 所有人在该时点均已知event-free
            brier[
                position
            ] = float(
                np.mean(
                    risk[
                        :,
                        position
                    ]
                    ** 2
                )
            )

        elif metric_time < test_max:

            _, values = brier_score(
                y_train,
                y_test,
                survival[
                    :,
                    position:
                    position + 1
                ],
                np.asarray(
                    [metric_time]
                ),
            )

            brier[
                position
            ] = float(
                values[0]
            )

        else:

            raise ValueError(
                "bootstrap样本不足以支持5年IBS。"
            )


    if not np.isfinite(
        brier
    ).all():

        raise ValueError(
            "Brier存在NaN。"
        )


    if hasattr(
        np,
        "trapezoid",
    ):

        integral = np.trapezoid(
            brier,
            METRIC_TIMES,
        )

    else:

        integral = np.trapz(
            brier,
            METRIC_TIMES,
        )


    return float(
        integral
        /
        (
            METRIC_TIMES[-1]
            -
            METRIC_TIMES[0]
        )
    )


# =============================================================================
# 12. ICI FUNCTIONS
# =============================================================================

def step_value(
    time_grid,
    probability,
    query,
):

    query = np.atleast_1d(
        np.asarray(
            query,
            dtype=float,
        )
    )

    index = (
        np.searchsorted(
            time_grid,
            query,
            side="right",
        )
        - 1
    )

    result = np.ones(
        len(query),
        dtype=float,
    )

    valid = (
        index >= 0
    )

    result[
        valid
    ] = probability[
        index[
            valid
        ]
    ]

    return result


def calculate_ici(
    predicted_risk,
    event,
    time,
    horizon=60.0,
):

    predicted_risk = np.asarray(
        predicted_risk,
        dtype=float,
    )

    event = np.asarray(
        event,
        dtype=bool,
    )

    time = np.asarray(
        time,
        dtype=float,
    )


    case = (
        event
        &
        (
            time <= horizon
        )
    )

    control = (
        (~event)
        &
        (
            time >= horizon
        )
    )

    known = (
        case
        |
        control
    )


    if case.sum() < 3:
        raise ValueError(
            "事件数过少，ICI不可稳定估计。"
        )


    censor_time, censor_survival = (
        kaplan_meier_estimator(
            event,
            time,
            reverse=True,
        )
    )


    weights = np.zeros(
        len(event),
        dtype=float,
    )


    if case.any():

        query = np.nextafter(
            time[
                case
            ],
            -np.inf,
        )

        g = step_value(
            censor_time,
            censor_survival,
            query,
        )

        weights[
            case
        ] = (
            1.0
            /
            np.clip(
                g,
                1e-6,
                None,
            )
        )


    if control.any():

        query = np.nextafter(
            horizon,
            -np.inf,
        )

        g = float(
            step_value(
                censor_time,
                censor_survival,
                query,
            )[0]
        )

        weights[
            control
        ] = (
            1.0
            /
            max(
                g,
                1e-6,
            )
        )


    p = np.clip(
        predicted_risk,
        1e-6,
        1 - 1e-6,
    )


    x = (
        np.log(
            p
            /
            (
                1
                - p
            )
        )
        .reshape(
            -1,
            1,
        )
    )


    y = case.astype(int)


    calibration_model = make_pipeline(

        SplineTransformer(
            n_knots=4,
            degree=3,
            include_bias=False,
        ),

        LogisticRegression(
            penalty="l2",
            C=1e6,
            solver="lbfgs",
            max_iter=3000,
        ),
    )


    calibration_model.fit(
        x[
            known
        ],
        y[
            known
        ],
        logisticregression__sample_weight=
            weights[
                known
            ],
    )


    fitted = (
        calibration_model
        .predict_proba(
            x
        )[
            :,
            1
        ]
    )


    return float(
        np.mean(
            np.abs(
                fitted
                -
                predicted_risk
            )
        )
    )


# =============================================================================
# 13. POINT ESTIMATES
# =============================================================================

point_ibs = {}
point_ici = {}


for (
    landmark_index,
    landmark_year
) in zip(
    PRIMARY_LANDMARK_INDEX,
    PRIMARY_LANDMARKS,
):

    temp = brier_point_df.loc[
        np.isclose(
            brier_point_df[
                "landmark_year"
            ],
            landmark_year,
        )
    ]

    point_ibs[
        landmark_year
    ] = float(
        temp[
            "IBS"
        ].iloc[0]
    )


    rows = np.where(
        landmark_index_long
        == landmark_index
    )[0]


    point_ici[
        landmark_year
    ] = calculate_ici(
        recal_risk_point[
            rows,
            9
        ],

        metadata.iloc[
            rows
        ][
            "event_within_60m"
        ]
        .to_numpy(int),

        metadata.iloc[
            rows
        ][
            "analysis_time_month"
        ]
        .to_numpy(float),
    )


# =============================================================================
# 14. BOOTSTRAP
# =============================================================================

rng = np.random.default_rng(
    RANDOM_SEED
)


boot_ibs = np.full(
    (
        BOOTSTRAP_REPS,
        4,
    ),
    np.nan,
)

boot_ici = np.full_like(
    boot_ibs,
    np.nan,
)


start_time = time.time()


for bootstrap_index in range(
    BOOTSTRAP_REPS
):

    # -------------------------------------------------------------------------
    # 患者级有放回抽样
    # -------------------------------------------------------------------------

    sampled_patients = rng.integers(
        0,
        N_PATIENT,
        size=N_PATIENT,
    )


    patient_counts = np.bincount(
        sampled_patients,
        minlength=N_PATIENT,
    )


    try:

        recal_risk_boot, recal_hazard_boot = (
            build_bootstrap_recalibrated_risk(
                patient_counts
            )
        )

    except Exception:

        continue


    for j, landmark_index in enumerate(
        PRIMARY_LANDMARK_INDEX
    ):

        original_rows = np.where(
            landmark_index_long
            == landmark_index
        )[0]


        # ---------------------------------------------------------------------
        # 每个patient-landmark row按患者bootstrap multiplicity重复
        # ---------------------------------------------------------------------

        counts = (
            patient_counts[
                patient_index_long[
                    original_rows
                ]
            ]
        )


        keep = (
            counts > 0
        )


        bootstrap_rows = np.repeat(
            original_rows[
                keep
            ],
            counts[
                keep
            ],
        )


        if len(
            bootstrap_rows
        ) < 20:

            continue


        # ---------------------------------------------------------------------
        # 相应OOF recalibrated prediction也重复
        # ---------------------------------------------------------------------

        bootstrap_risk = (
            recal_risk_boot[
                bootstrap_rows
            ]
        )


        if not np.isfinite(
            bootstrap_risk
        ).all():

            continue


        try:

            boot_ibs[
                bootstrap_index,
                j
            ] = calculate_ibs(
                landmark_index,
                bootstrap_rows,
                bootstrap_risk,
            )


            boot_ici[
                bootstrap_index,
                j
            ] = calculate_ici(
                bootstrap_risk[
                    :,
                    9
                ],

                metadata.iloc[
                    bootstrap_rows
                ][
                    "event_within_60m"
                ]
                .to_numpy(int),

                metadata.iloc[
                    bootstrap_rows
                ][
                    "analysis_time_month"
                ]
                .to_numpy(float),
            )


        except Exception:

            continue


    if (
        bootstrap_index + 1
    ) % 50 == 0:

        elapsed = (
            time.time()
            - start_time
        ) / 60.0

        print(
            f"Bootstrap "
            f"{bootstrap_index + 1}/"
            f"{BOOTSTRAP_REPS} | "
            f"elapsed={elapsed:.1f} min"
        )


# =============================================================================
# 15. SAVE BOOTSTRAP
# =============================================================================

np.savez_compressed(
    FINAL_DIR
    / "IBS_ICI_crossfit_bootstrap_1000.npz",

    IBS=boot_ibs,

    ICI=boot_ici,

    landmarks=
        PRIMARY_LANDMARKS,

    bootstrap_reps=
        BOOTSTRAP_REPS,

    random_seed=
        RANDOM_SEED,
)


# =============================================================================
# 16. CI
# =============================================================================

def percentile_ci(
    values
):

    values = np.asarray(
        values,
        dtype=float,
    )

    values = values[
        np.isfinite(
            values
        )
    ]

    if len(values) == 0:
        return (
            np.nan,
            np.nan,
            0,
        )


    return (
        float(
            np.quantile(
                values,
                0.025,
            )
        ),

        float(
            np.quantile(
                values,
                0.975,
            )
        ),

        len(values),
    )


ci_records = []


for j, landmark_year in enumerate(
    PRIMARY_LANDMARKS
):

    ibs_lower, ibs_upper, ibs_n = (
        percentile_ci(
            boot_ibs[
                :,
                j
            ]
        )
    )


    ici_lower, ici_upper, ici_n = (
        percentile_ci(
            boot_ici[
                :,
                j
            ]
        )
    )


    ci_records.append({

        "landmark_year":
            landmark_year,

        "IBS":
            point_ibs[
                landmark_year
            ],

        "IBS_lower_95":
            ibs_lower,

        "IBS_upper_95":
            ibs_upper,

        "IBS_valid_bootstrap_n":
            ibs_n,

        "ICI_5y":
            point_ici[
                landmark_year
            ],

        "ICI_lower_95":
            ici_lower,

        "ICI_upper_95":
            ici_upper,

        "ICI_valid_bootstrap_n":
            ici_n,
    })


ci_df = pd.DataFrame(
    ci_records
)


ci_df.to_csv(
    FINAL_DIR
    / "IBS_ICI_95CI_summary.csv",
    index=False,
    encoding="utf-8-sig",
)


# =============================================================================
# 17. VALIDITY CHECK
# =============================================================================

print()
print(
    "Bootstrap有效次数："
)

print(
    ci_df[
        [
            "landmark_year",
            "IBS_valid_bootstrap_n",
            "ICI_valid_bootstrap_n",
        ]
    ].to_string(
        index=False
    )
)


if (
    ci_df[
        [
            "IBS_valid_bootstrap_n",
            "ICI_valid_bootstrap_n",
        ]
    ].min().min()
    < 800
):

    warnings.warn(
        "至少一个指标有效bootstrap少于800次，"
        "需要进一步检查。"
    )


# =============================================================================
# 18. THREE SIGNIFICANT DIGITS
# =============================================================================

def format_sig(
    value,
    sig=3
):

    value = float(value)


    if not np.isfinite(value):

        return ""


    if value == 0:

        return "0"


    magnitude = int(
        np.floor(
            np.log10(
                abs(value)
            )
        )
    )


    decimals = (
        sig
        - magnitude
        - 1
    )


    if decimals > 0:

        return f"{value:.{decimals}f}"


    return f"{value:.0f}"


def format_ci(
    point,
    lower,
    upper
):

    return (
        f"{format_sig(point)}"
        f"（{format_sig(lower)}，"
        f"{format_sig(upper)}）"
    )


# =============================================================================
# 19. BUILD FINAL CORE TABLE
# =============================================================================

results = {}


for landmark_year in (
    PRIMARY_LANDMARKS
):

    perf = landmark_df.loc[
        np.isclose(
            landmark_df[
                "landmark_year"
            ],
            landmark_year,
        )
    ].iloc[0]


    ci = ci_df.loc[
        np.isclose(
            ci_df[
                "landmark_year"
            ],
            landmark_year,
        )
    ].iloc[0]


    column = COLUMN_NAMES[
        landmark_year
    ]


    results[
        column
    ] = {

        "Patients at risk, n":
            str(
                int(
                    perf[
                        "risk_set_n"
                    ]
                )
            ),


        "CKD events within 5 years, n":
            str(
                int(
                    perf[
                        "event_within_5y_n"
                    ]
                )
            ),


        "5-year Uno C-index (95% CI)":
            format_ci(
                perf[
                    "uno_c_index_5y"
                ],
                perf[
                    "uno_c_lower_95"
                ],
                perf[
                    "uno_c_upper_95"
                ],
            ),


        "iAUC (95% CI)":
            format_ci(
                perf[
                    "integrated_dynamic_auc"
                ],
                perf[
                    "iauc_lower_95"
                ],
                perf[
                    "iauc_upper_95"
                ],
            ),


        "IBS (95% CI)":
            format_ci(
                ci[
                    "IBS"
                ],
                ci[
                    "IBS_lower_95"
                ],
                ci[
                    "IBS_upper_95"
                ],
            ),


        "5-year ICI (95% CI)":
            format_ci(
                ci[
                    "ICI_5y"
                ],
                ci[
                    "ICI_lower_95"
                ],
                ci[
                    "ICI_upper_95"
                ],
            ),
    }


ROW_ORDER = [

    "Patients at risk, n",

    "CKD events within 5 years, n",

    "5-year Uno C-index (95% CI)",

    "iAUC (95% CI)",

    "IBS (95% CI)",

    "5-year ICI (95% CI)",
]


final_table = pd.DataFrame(
    results
).loc[
    ROW_ORDER
]


final_table.index.name = (
    "Performance metric"
)


# =============================================================================
# 20. SAVE FINAL TABLE
# =============================================================================

OUTPUT_CSV = (
    FINAL_DIR
    / "Table1_External_validation_FINAL_core_metrics_with_95CI.csv"
)

OUTPUT_XLSX = (
    FINAL_DIR
    / "Table1_External_validation_FINAL_core_metrics_with_95CI.xlsx"
)


final_table.to_csv(
    OUTPUT_CSV,
    encoding="utf-8-sig",
)


final_table.to_excel(
    OUTPUT_XLSX,
    sheet_name="External validation",
)


# =============================================================================
# 21. FORMAT EXCEL
# =============================================================================

wb = load_workbook(
    OUTPUT_XLSX
)

ws = wb[
    "External validation"
]


thin = Side(
    style="thin",
    color="000000",
)


for row in ws.iter_rows():

    for cell in row:

        cell.font = Font(
            name="Arial",
            size=11,
            color="000000",
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

        cell.border = Border(
            bottom=thin,
        )


for cell in ws[1]:

    cell.font = Font(
        name="Arial",
        size=11,
        bold=True,
        color="000000",
    )


for row in range(
    2,
    ws.max_row + 1
):

    ws.cell(
        row=row,
        column=1,
    ).alignment = Alignment(
        horizontal="left",
        vertical="center",
    )


ws.column_dimensions[
    "A"
].width = 38


for column in [
    "B",
    "C",
    "D",
    "E",
]:

    ws.column_dimensions[
        column
    ].width = 29


for row in range(
    1,
    ws.max_row + 1
):

    ws.row_dimensions[
        row
    ].height = 24


ws.freeze_panes = "B2"


wb.save(
    OUTPUT_XLSX
)


# =============================================================================
# 22. PRINT FINAL TABLE
# =============================================================================

print()
print("=" * 140)

print(
    "Table. Geographic external validation of "
    "landmark-specific dynamic LSTM predictions for incident CKD"
)

print("=" * 140)


print(
    final_table.to_string()
)


print()
print(
    "最终CSV："
)

print(
    OUTPUT_CSV
)


print()
print(
    "最终Excel："
)

print(
    OUTPUT_XLSX
)


print()
print(
    "IBS / ICI bootstrap summary："
)

print(
    FINAL_DIR
    / "IBS_ICI_95CI_summary.csv"
)


print("=" * 140)