from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.lines import Line2D
from matplotlib.patches import Patch
import openpyxl


SOURCE = Path(r"__CKD_LOCAL_INPUT__/Users/Administrator/Desktop/HIV患者随访数据最初.xlsx")
ORIGINAL_SOURCE = Path(r"__CKD_LOCAL_INPUT__/多中心艾滋ML/NC修稿数据/深圳南宁随访数据表.xlsx")
OUT_DIR = Path(r"__CKD_LOCAL_INPUT__/Documents/深度学习/output/pdf")
TMP_DIR = Path(r"__CKD_LOCAL_INPUT__/Documents/深度学习/tmp/pdfs")
OUT_PDF = OUT_DIR / "Figure_S_Preimputation_Missingness_Pareto_with_excluded_variables.pdf"
QA_PNG = TMP_DIR / "Figure_S_Preimputation_Missingness_Pareto_with_excluded_variables.png"

LAB_VARIABLES = [
    "HIVRNA", "CD4", "CD8", "SCr", "Urea", "WBC", "PLT", "HB",
    "TC", "TG", "HDL", "LDL", "GLU", "ALT", "AST", "eGFR",
]
EXCLUDED_HIGH_MISSING_VARIABLES = ["HbA1c", "CRP", "UACR", "P", "Ca", "K", "Na"]
FINAL_MODEL_VARIABLES = {
    "HIVRNA", "CD4", "CD8", "Urea", "WBC", "PLT", "HB", "TC",
    "TG", "HDL", "LDL", "GLU", "ALT", "AST", "eGFR",
}


def is_missing(value):
    return value is None or (isinstance(value, str) and not value.strip())


workbook = openpyxl.load_workbook(SOURCE, read_only=True, data_only=True)
sheet = workbook["Sheet1"]
headers = list(next(sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
indices = {name: headers.index(name) for name in LAB_VARIABLES}
id_index = headers.index("ID")
time_index = headers.index("time_bin")

missing_counts = {name: 0 for name in LAB_VARIABLES}
n_observations = 0
for row in sheet.iter_rows(min_row=2, values_only=True):
    if is_missing(row[id_index]) or is_missing(row[time_index]):
        continue
    n_observations += 1
    for name, index in indices.items():
        if is_missing(row[index]):
            missing_counts[name] += 1
workbook.close()

original_workbook = openpyxl.load_workbook(ORIGINAL_SOURCE, read_only=True, data_only=True)
original_sheet = original_workbook["Sheet1"]
original_headers = list(next(original_sheet.iter_rows(min_row=1, max_row=1, values_only=True)))
original_indices = {name: original_headers.index(name) for name in EXCLUDED_HIGH_MISSING_VARIABLES}
original_id_index = original_headers.index("ID")
original_time_index = original_headers.index("time_bin")
original_missing_counts = {name: 0 for name in EXCLUDED_HIGH_MISSING_VARIABLES}
n_original_observations = 0
for row in original_sheet.iter_rows(min_row=2, values_only=True):
    if is_missing(row[original_id_index]) or is_missing(row[original_time_index]):
        continue
    n_original_observations += 1
    for name, index in original_indices.items():
        if is_missing(row[index]):
            original_missing_counts[name] += 1
original_workbook.close()
if n_original_observations != n_observations:
    raise ValueError(
        f"Aligned observation counts differ: current={n_observations}, original={n_original_observations}"
    )

rows = [
    {
        "variable": name,
        "display": "HIV RNA" if name == "HIVRNA" else name,
        "missing_n": missing_counts[name],
        "missing_pct": missing_counts[name] / n_observations * 100,
        "included": name in FINAL_MODEL_VARIABLES,
    }
    for name in LAB_VARIABLES
]
rows.extend(
    {
        "variable": name,
        "display": name,
        "missing_n": original_missing_counts[name],
        "missing_pct": original_missing_counts[name] / n_original_observations * 100,
        "included": False,
    }
    for name in EXCLUDED_HIGH_MISSING_VARIABLES
)
rows.sort(key=lambda item: (-item["missing_pct"], item["display"].lower()))

total_missing = sum(item["missing_n"] for item in rows)
running = 0
cumulative = []
for item in rows:
    running += item["missing_n"]
    cumulative.append(running / total_missing * 100)

plt.rcParams.update({"font.family": "Arial", "pdf.fonttype": 42, "ps.fonttype": 42})
fig, ax = plt.subplots(figsize=(13.5, 7.2))
x = list(range(len(rows)))
bars = ax.bar(
    x,
    [item["missing_n"] for item in rows],
    width=0.82,
    color=["#E53924" if item["missing_pct"] > 25 else "#FFF3A6" for item in rows],
    edgecolor=["#000000" if item["included"] else "#8C8C8C" for item in rows],
    linewidth=[1.6 if item["included"] else 0.8 for item in rows],
)
ax.set_xticks(x, [item["display"] for item in rows], rotation=45, ha="right")
ax.set_ylabel("Frequency of missing values", fontsize=10)
ax.set_xlabel("Continuous laboratory predictors", fontsize=10)
ax.grid(axis="y", color="#D9D9D9", linewidth=0.6)
ax.set_axisbelow(True)
ax.tick_params(axis="x", labelsize=8.5)
ax.tick_params(axis="y", labelsize=8.5)
for spine in ("top", "right"):
    ax.spines[spine].set_visible(False)

offset = max(item["missing_n"] for item in rows) * 0.015
for bar, item in zip(bars, rows):
    ax.text(
        bar.get_x() + bar.get_width() / 2,
        bar.get_height() + offset,
        f'{item["missing_pct"]:.1f}%',
        ha="center",
        va="bottom",
        fontsize=7.5,
        color="#333333",
    )

ax2 = ax.twinx()
ax2.plot(x, cumulative, color="#111111", marker="o", markersize=4.2, linewidth=1.5)
ax2.set_ylabel("Cumulative share of missing values (%)", fontsize=10)
ax2.set_ylim(0, 105)
ax2.set_yticks([0, 25, 50, 75, 100])
ax2.tick_params(axis="y", labelsize=8.5)
ax2.spines["top"].set_visible(False)

fig.legend(
    handles=[
        Patch(facecolor="#E53924", edgecolor="#666666", label=">25% missing"),
        Patch(facecolor="#FFF3A6", edgecolor="#666666", label="≤25% missing"),
        Patch(facecolor="white", edgecolor="#000000", linewidth=1.6, label="Included in final model"),
        Line2D([0], [0], color="#111111", marker="o", linewidth=1.5, label="Cumulative share"),
    ],
    loc="upper center",
    bbox_to_anchor=(0.53, 0.965),
    frameon=False,
    ncol=4,
    fontsize=8.5,
)
fig.text(
    0.08,
    0.025,
    "Sources: current follow-up workbook for retained variables and original candidate-variable workbook for excluded high-missing variables. Missingness used 397,410 aligned patient-time observations before imputation.",
    ha="left",
    va="bottom",
    fontsize=7.5,
    color="#4D4D4D",
)
fig.subplots_adjust(left=0.08, right=0.92, top=0.84, bottom=0.24)

OUT_DIR.mkdir(parents=True, exist_ok=True)
TMP_DIR.mkdir(parents=True, exist_ok=True)
fig.savefig(OUT_PDF, format="pdf", bbox_inches="tight", facecolor="white")
fig.savefig(QA_PNG, dpi=220, bbox_inches="tight", facecolor="white")
plt.close(fig)

print(f"observations={n_observations}")
for item in rows:
    print(f'{item["variable"]}\t{item["missing_n"]}\t{item["missing_pct"]:.6f}\t{item["included"]}')
print(OUT_PDF)
