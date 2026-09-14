from __future__ import annotations

from dataclasses import dataclass
from hashlib import sha256
from math import erfc, exp, lgamma, log, sqrt
from pathlib import Path
import shutil

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


CSV_PATH = Path(r"__CKD_LOCAL_INPUT__/多中心艾滋ML/NC修稿数据/深圳南宁随访数据表2.csv")
REFERENCE = Path(
    r"__CKD_LOCAL_INPUT__/Documents/深度学习/tmp/docx/baseline_development/Table1_baseline_reference.docx"
)
FINAL = Path(
    r"__CKD_LOCAL_INPUT__/多中心艾滋ML/CKD_ML_投稿/NC/Table1_baseline_development_updated.docx"
)
AUDIT_CSV = Path(
    r"__CKD_LOCAL_INPUT__/Documents/深度学习/tmp/docx/baseline_development/Table1_baseline_development_statistics.csv"
)
EXPECTED_REFERENCE_SHA256 = (
    "3EF413DBA9F336DBFDDB43011064D88B7893EF7BC769EA31911E70AAB32FD0F5"
)

FONT = "Times New Roman"
FONT_SIZE = Pt(9.5)
GRID_DXA = [2593, 2293, 2322, 2293, 946]


def file_sha256(path: Path) -> str:
    digest = sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest().upper()


def load_baseline() -> pd.DataFrame:
    pieces: list[pd.DataFrame] = []
    for chunk in pd.read_csv(
        CSV_PATH, encoding="gb18030", chunksize=100_000, low_memory=False
    ):
        pieces.append(chunk.loc[chunk["time_bin"].eq(0)].copy())
    df = pd.concat(pieces, ignore_index=True)

    if df["ID"].duplicated().any():
        raise ValueError("time_bin=0 contains duplicate participant IDs")
    if not set(df["CKDstatus"].unique()).issubset({0, 1}):
        raise ValueError("CKDstatus must be binary")

    df["Sex_table"] = (
        df["Sex"].astype(str).str.strip().str.lower().map(
            {"female": "Female", "male": "Male"}
        )
    )
    df["Site_table"] = df["data"].map({"深圳": "Shenzhen", "南宁": "Nanning"})
    df["Marriage_table"] = df["Marriage"].replace(
        {
            "Divorced-separated-or-widowed": "Divorced, separated, or widowed",
        }
    )
    df["Course_table"] = df["Course"].replace(
        {
            "Heterosexual": "Heterosexual contact",
            "Male to male": "Male-to-male sexual contact",
            "Drugs": "Injecting drug use",
        }
    )
    df["Oppinfection_table"] = df["Oppinfection"].map({0: "No", 1: "Yes"})
    df["WHOstage_table"] = df["WHOstage"].map(
        {1: "I", 2: "II", 3: "III", 4: "IV"}
    )

    binary_columns = {
        "CVD_table": "CVD_status",
        "Diabetes_table": "diabetes_status",
        "Hypertension_table": "hypertension_status",
        "Hypercholesterolemia_table": "hypercholesterolemia_status",
        "Antidiabetic_table": "antidiabetic_med",
        "Antihypertensive_table": "antihypertensive_med",
        "Antilipid_table": "antilipid_med",
        "HBV_table": "HBV_status",
        "HCV_table": "HCV_status",
    }
    for target, source in binary_columns.items():
        df[target] = df[source].map({0: "No", 1: "Yes"})

    df["CD4_CD8_ratio"] = df["CD4"] / df["CD8"]

    regimen_map = [
        ("current_TDF_NNRTI_3TC_FTC", "TDF + NNRTI + 3TC/FTC"),
        ("current_TDF_PI_3TC_FTC", "TDF + PI + 3TC/FTC"),
        ("current_nonTDF_PI", "Non-TDF PI-based"),
        ("current_BIC_FTC_TAF", "BIC/FTC/TAF"),
        ("current_EVGc_FTC_TAF", "EVG/c/FTC/TAF"),
        ("current_TDF_INSTI_3TC_FTC", "TDF + INSTI + 3TC/FTC"),
        ("current_nonTDF_DTG", "Non-TDF DTG-based"),
        (
            "current_nonTDF_traditional_NNRTI",
            "Non-TDF traditional NNRTI-based",
        ),
    ]
    regimen_cols = [item[0] for item in regimen_map]
    regimen_sum = df[regimen_cols].sum(axis=1)
    if (regimen_sum > 1).any():
        raise ValueError("More than one current ART regimen indicator is active")
    df["ART_regimen_table"] = "Other/unspecified"
    for source, label in regimen_map:
        df.loc[df[source].eq(1), "ART_regimen_table"] = label

    required_table_columns = [
        "Sex_table",
        "Site_table",
        "Marriage_table",
        "Course_table",
        "Oppinfection_table",
        "WHOstage_table",
        *binary_columns.keys(),
        "ART_regimen_table",
        "CD4_CD8_ratio",
    ]
    missing_required = {
        col: int(df[col].isna().sum())
        for col in required_table_columns
        if df[col].isna().any()
    }
    if missing_required:
        raise ValueError(f"Unmapped or missing table values: {missing_required}")
    return df


def gamma_q(a: float, x: float) -> float:
    """Regularized upper incomplete gamma Q(a, x)."""
    if x < 0 or a <= 0:
        return float("nan")
    if x == 0:
        return 1.0
    eps = 3e-14
    fpmin = 1e-300
    itmax = 1000
    gln = lgamma(a)
    if x < a + 1.0:
        ap = a
        summation = 1.0 / a
        delta = summation
        for _ in range(itmax):
            ap += 1.0
            delta *= x / ap
            summation += delta
            if abs(delta) < abs(summation) * eps:
                break
        p = summation * exp(-x + a * log(x) - gln)
        return max(0.0, min(1.0, 1.0 - p))

    b = x + 1.0 - a
    c = 1.0 / fpmin
    d = 1.0 / b
    h = d
    for i in range(1, itmax + 1):
        an = -i * (i - a)
        b += 2.0
        d = an * d + b
        if abs(d) < fpmin:
            d = fpmin
        c = b + an / c
        if abs(c) < fpmin:
            c = fpmin
        d = 1.0 / d
        delta = d * c
        h *= delta
        if abs(delta - 1.0) < eps:
            break
    q = exp(-x + a * log(x) - gln) * h
    return max(0.0, min(1.0, q))


def fisher_two_sided(table: np.ndarray) -> float:
    a, b, c, d = [int(v) for v in table.ravel()]
    r1, r2 = a + b, c + d
    c1 = a + c
    n = r1 + r2
    lo = max(0, c1 - r2)
    hi = min(r1, c1)

    def log_choose(n_: int, k_: int) -> float:
        return lgamma(n_ + 1) - lgamma(k_ + 1) - lgamma(n_ - k_ + 1)

    def probability(a_: int) -> float:
        return exp(log_choose(r1, a_) + log_choose(r2, c1 - a_) - log_choose(n, c1))

    observed = probability(a)
    p = sum(probability(x) for x in range(lo, hi + 1) if probability(x) <= observed + 1e-15)
    return min(1.0, p)


def categorical_p(series: pd.Series, outcome: pd.Series) -> tuple[float, str]:
    table = pd.crosstab(series, outcome).reindex(columns=[0, 1], fill_value=0)
    observed = table.to_numpy(dtype=float)
    if observed.shape[0] < 2 or observed.sum() == 0:
        return float("nan"), "not testable"
    expected = observed.sum(axis=1, keepdims=True) @ observed.sum(axis=0, keepdims=True) / observed.sum()
    if observed.shape == (2, 2) and (expected < 5).any():
        return fisher_two_sided(observed.astype(int)), "Fisher exact"
    statistic = float(((observed - expected) ** 2 / expected).sum())
    degrees_freedom = (observed.shape[0] - 1) * (observed.shape[1] - 1)
    return gamma_q(degrees_freedom / 2.0, statistic / 2.0), "Pearson chi-square"


def mann_whitney_p(x: pd.Series, y: pd.Series) -> float:
    x = pd.to_numeric(x, errors="coerce").dropna().to_numpy(dtype=float)
    y = pd.to_numeric(y, errors="coerce").dropna().to_numpy(dtype=float)
    n1, n2 = len(x), len(y)
    pooled = np.concatenate([x, y])
    ranks = pd.Series(pooled).rank(method="average").to_numpy()
    u1 = ranks[:n1].sum() - n1 * (n1 + 1) / 2.0
    mean_u = n1 * n2 / 2.0
    _, tie_counts = np.unique(pooled, return_counts=True)
    n = n1 + n2
    tie_term = float(((tie_counts**3) - tie_counts).sum())
    variance = n1 * n2 / 12.0 * ((n + 1.0) - tie_term / (n * (n - 1.0)))
    if variance <= 0:
        return 1.0
    z = max(0.0, abs(u1 - mean_u) - 0.5) / sqrt(variance)
    return erfc(z / sqrt(2.0))


def format_p(p: float) -> str:
    if not np.isfinite(p):
        return "NA"
    if p < 0.001:
        return "<0.001"
    return f"{p:.3f}"


def continuous_summary(series: pd.Series) -> str:
    values = pd.to_numeric(series, errors="coerce").dropna()
    return f"{values.median():.2f} [{values.quantile(0.25):.2f}, {values.quantile(0.75):.2f}]"


def categorical_summary(series: pd.Series, level: str) -> str:
    numerator = int(series.eq(level).sum())
    denominator = int(series.notna().sum())
    return f"{numerator:,} ({100 * numerator / denominator:.1f})"


@dataclass
class TableRow:
    label: str
    overall: str = ""
    non_ckd: str = ""
    ckd: str = ""
    p_value: str = ""
    role: str = "category"  # header, continuous, group, category
    test: str = ""
    source_variable: str = ""


def assemble_rows(df: pd.DataFrame) -> list[TableRow]:
    rows: list[TableRow] = []
    n_all = len(df)
    n_non = int(df["CKDstatus"].eq(0).sum())
    n_ckd = int(df["CKDstatus"].eq(1).sum())
    rows.append(
        TableRow(
            "Characteristic",
            f"Overall\n(n = {n_all:,})",
            f"Non-CKD\n(n = {n_non:,})",
            f"CKD\n(n = {n_ckd:,})",
            "P value",
            role="header",
        )
    )

    non = df.loc[df["CKDstatus"].eq(0)]
    ckd = df.loc[df["CKDstatus"].eq(1)]

    def add_continuous(label: str, variable: str) -> None:
        p = mann_whitney_p(non[variable], ckd[variable])
        rows.append(
            TableRow(
                label,
                continuous_summary(df[variable]),
                continuous_summary(non[variable]),
                continuous_summary(ckd[variable]),
                format_p(p),
                role="continuous",
                test="Mann-Whitney U",
                source_variable=variable,
            )
        )

    def add_categorical(label: str, variable: str, levels: list[str]) -> None:
        p, test = categorical_p(df[variable], df["CKDstatus"])
        rows.append(TableRow(label, role="group", test=test, source_variable=variable))
        for index, level in enumerate(levels):
            rows.append(
                TableRow(
                    level,
                    categorical_summary(df[variable], level),
                    categorical_summary(non[variable], level),
                    categorical_summary(ckd[variable], level),
                    format_p(p) if index == 0 else "",
                    role="category",
                    test=test if index == 0 else "",
                    source_variable=variable,
                )
            )

    add_continuous("Age, years", "Age")
    add_categorical("Study center", "Site_table", ["Shenzhen", "Nanning"])
    add_categorical("Sex", "Sex_table", ["Female", "Male"])
    add_categorical(
        "Marriage status",
        "Marriage_table",
        [
            "Never married",
            "Married or cohabiting",
            "Divorced, separated, or widowed",
            "Others",
        ],
    )
    add_categorical(
        "HIV transmission route",
        "Course_table",
        [
            "Heterosexual contact",
            "Male-to-male sexual contact",
            "Injecting drug use",
            "Others",
        ],
    )
    add_categorical("Opportunistic infections", "Oppinfection_table", ["No", "Yes"])
    add_categorical("WHO stage", "WHOstage_table", ["I", "II", "III", "IV"])

    for label, variable in [
        ("BMI, kg/m²", "BMI"),
        ("HIV RNA, log10 copies/mL", "HIVRNA_log10"),
        ("CD4 count, cells/μL", "CD4"),
        ("CD8 count, cells/μL", "CD8"),
        ("CD4/CD8 ratio", "CD4_CD8_ratio"),
        ("Serum creatinine, μmol/L", "SCr"),
        ("eGFR, mL/min/1.73m²", "eGFR"),
        ("Urea, mmol/L", "Urea"),
        ("WBC, 10^9/L", "WBC"),
        ("PLT, 10^9/L", "PLT"),
        ("HB, g/L", "HB"),
        ("TC, mmol/L", "TC"),
        ("TG, mmol/L", "TG"),
        ("HDL, mmol/L", "HDL"),
        ("LDL, mmol/L", "LDL"),
        ("Glucose, mmol/L", "GLU"),
        ("ALT, U/L", "ALT"),
        ("AST, U/L", "AST"),
    ]:
        add_continuous(label, variable)

    for label, variable in [
        ("CVD", "CVD_table"),
        ("Hypertension", "Hypertension_table"),
        ("Diabetes", "Diabetes_table"),
        ("Hypercholesterolemia", "Hypercholesterolemia_table"),
        ("Antihypertensive medication", "Antihypertensive_table"),
        ("Antidiabetic medication", "Antidiabetic_table"),
        ("Lipid-lowering medication", "Antilipid_table"),
        ("HBV", "HBV_table"),
        ("HCV", "HCV_table"),
    ]:
        add_categorical(label, variable, ["No", "Yes"])

    add_categorical(
        "Current ART regimen",
        "ART_regimen_table",
        [
            "TDF + NNRTI + 3TC/FTC",
            "TDF + PI + 3TC/FTC",
            "Non-TDF PI-based",
            "BIC/FTC/TAF",
            "EVG/c/FTC/TAF",
            "TDF + INSTI + 3TC/FTC",
            "Non-TDF DTG-based",
            "Non-TDF traditional NNRTI-based",
            "Other/unspecified",
        ],
    )
    return rows


def set_run_font(run, *, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = FONT
    run.font.size = FONT_SIZE
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), FONT)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), FONT)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def clear_paragraph(paragraph) -> None:
    p = paragraph._p
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def remove_extra_paragraphs(cell) -> None:
    while len(cell.paragraphs) > 1:
        cell._tc.remove(cell.paragraphs[-1]._p)


def write_cell(cell, text: str, *, role: str, column: int) -> None:
    remove_extra_paragraphs(cell)
    paragraph = cell.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    paragraph.paragraph_format.keep_with_next = role == "group"
    paragraph.paragraph_format.widow_control = True
    if role == "category" and column == 0:
        paragraph.paragraph_format.left_indent = Pt(10)
        paragraph.paragraph_format.first_line_indent = Pt(0)
    else:
        paragraph.paragraph_format.left_indent = Pt(0)
        paragraph.paragraph_format.first_line_indent = Pt(0)

    if role == "header" and column == 4:
        run = paragraph.add_run("P")
        set_run_font(run, bold=True, italic=True)
        run = paragraph.add_run(" value")
        set_run_font(run, bold=True)
    elif "\n" in text:
        first, second = text.split("\n", 1)
        run = paragraph.add_run(first)
        set_run_font(run, bold=role == "header")
        run.add_break()
        run = paragraph.add_run(second)
        set_run_font(run, bold=role == "header")
    else:
        run = paragraph.add_run(text)
        bold = role == "header" or (column == 0 and role in {"continuous", "group"})
        set_run_font(run, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_border(cell, *, top: bool = False, bottom: bool = False) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    old = tc_pr.find(qn("w:tcBorders"))
    if old is not None:
        tc_pr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge, active in [("top", top), ("left", False), ("bottom", bottom), ("right", False)]:
        element = OxmlElement(f"w:{edge}")
        if active:
            element.set(qn("w:val"), "single")
            element.set(qn("w:color"), "auto")
            element.set(qn("w:sz"), "8")
            element.set(qn("w:space"), "0")
        else:
            element.set(qn("w:val"), "nil")
        borders.append(element)
    tc_pr.append(borders)


def set_table_geometry(table) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(GRID_DXA)))
    tbl_w.set(qn("w:type"), "dxa")

    grid_cols = table._tbl.tblGrid.gridCol_lst
    for grid_col, width in zip(grid_cols, GRID_DXA):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, GRID_DXA):
            set_cell_width(cell, width)


def set_title(document: Document) -> None:
    paragraph = document.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(
        "Table 1. Baseline characteristics of participants in the development cohort"
    )
    set_run_font(run, bold=True)
    run.font.size = Pt(10)


def set_notes(document: Document) -> None:
    paragraph = document.paragraphs[-1]
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    note = (
        "Notes: Data are n (%) or median [interquartile range]. P values used Mann-Whitney U or chi-square/Fisher "
        "exact tests. Statistics use time_bin = 0 records after revised imputation; no displayed covariate was "
        "missing. CD4/CD8 was derived; eGFR used the 2021 CKD-EPI equation. SCr, serum creatinine; CVD, cardiovascular "
        "disease; HBV/HCV, hepatitis B/C virus; ART, antiretroviral therapy; TDF/TAF, tenofovir disoproxil "
        "fumarate/alafenamide; 3TC/FTC, lamivudine/emtricitabine; BIC/DTG/EVG/c, bictegravir/dolutegravir/"
        "elvitegravir-cobicistat; NNRTI/PI/INSTI, non-nucleoside reverse transcriptase/protease/integrase inhibitors."
    )
    run = paragraph.add_run(note)
    set_run_font(run)
    run.font.size = Pt(7.5)


def build_document(rows: list[TableRow]) -> None:
    if file_sha256(REFERENCE) != EXPECTED_REFERENCE_SHA256:
        raise RuntimeError("Reference copy no longer matches the distilled template")
    FINAL.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(REFERENCE, FINAL)
    document = Document(FINAL)
    set_title(document)

    table = document.tables[0]
    while len(table.rows) < len(rows):
        table.add_row()
    while len(table.rows) > len(rows):
        table._tbl.remove(table.rows[-1]._tr)

    for row_index, item in enumerate(rows):
        values = [item.label, item.overall, item.non_ckd, item.ckd, item.p_value]
        for column_index, (cell, value) in enumerate(zip(table.rows[row_index].cells, values)):
            write_cell(cell, value, role=item.role, column=column_index)
            set_cell_border(
                cell,
                top=row_index == 0,
                bottom=row_index == 0 or row_index == len(rows) - 1,
            )

    set_table_geometry(table)
    set_notes(document)
    document.save(FINAL)


def write_audit(rows: list[TableRow]) -> None:
    pd.DataFrame(
        [
            {
                "characteristic": row.label,
                "overall": row.overall.replace("\n", " "),
                "non_ckd": row.non_ckd.replace("\n", " "),
                "ckd": row.ckd.replace("\n", " "),
                "p_value": row.p_value,
                "row_role": row.role,
                "test": row.test,
                "source_variable": row.source_variable,
            }
            for row in rows
        ]
    ).to_csv(AUDIT_CSV, index=False, encoding="utf-8-sig")


def main() -> None:
    df = load_baseline()
    rows = assemble_rows(df)
    write_audit(rows)
    build_document(rows)
    print(f"participants={len(df):,}")
    print(f"non_ckd={df['CKDstatus'].eq(0).sum():,}")
    print(f"ckd={df['CKDstatus'].eq(1).sum():,}")
    print(f"table_rows={len(rows)}")
    print(f"output={FINAL}")


if __name__ == "__main__":
    main()
