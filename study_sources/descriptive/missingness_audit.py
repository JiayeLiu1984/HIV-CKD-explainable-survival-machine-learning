import csv
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook


SOURCE = Path(r"__CKD_LOCAL_INPUT__/多中心艾滋ML/NC修稿数据/深圳南宁随访数据表.xlsx")
FEATURE_GROUPS = Path(r"__CKD_WORKDIR__\rolling_5y_finaldata_step3_raw_features\feature_groups.json")
RNN_FEATURES = Path(r"__CKD_WORKDIR__\rolling_5y_step9b_rnn_tune_resume_v1\final_oof\rnn_feature_partition.csv")
LSTM_FEATURES = Path(r"__CKD_WORKDIR__\rolling_5y_finaldata_step10e_lstm_v2_final_oof_selected_medium12\lstm_v2_feature_partition.csv")
COX_FEATURES = Path(r"__CKD_WORKDIR__\rolling_5y_step7_unpenalized_cox_fixed95_v4\cox_fixed_feature_dictionary_95.csv")
RSF_FEATURES = Path(r"__CKD_WORKDIR__\rolling_5y_step8_pooled_landmark_rsf_resume_v2\rsf_feature_dictionary_96.csv")
OUT_DIR = Path(r"__CKD_LOCAL_INPUT__/Documents/深度学习/tmp/missingness_audit_20260905")
OUT_CSV = OUT_DIR / "candidate_predictor_missingness_audit.csv"
OUT_JSON = OUT_DIR / "candidate_predictor_missingness_summary.json"

NON_PREDICTORS = {"ID", "data", "time_bin", "month", "CKDstatus", "CKDtime", "interval"}
STATIC_VARS = {"Sex", "Age", "Marriage", "Course", "Oppinfection", "WHOstage", "BMI"}
EXCLUDED_HIGH_MISSING = {"Na", "P", "Ca", "K", "CRP", "HbA1c", "UACR"}
MISSING_STRINGS = {"", "na", "n/a", "nan", "null", "none", "."}


def is_missing(value):
    if value is None:
        return True
    if isinstance(value, str) and value.strip().lower() in MISSING_STRINGS:
        return True
    return False


def normalize_source_name(name):
    if name == "HIVRNA_log10":
        return "HIVRNA"
    for prefix, source in (
        ("Sex_", "Sex"),
        ("Marriage_", "Marriage"),
        ("Course_", "Course"),
        ("WHOstage_", "WHOstage"),
    ):
        if name.startswith(prefix):
            return source
    for suffix in ("_observed", "_time_since_last", "_delta_last_observed"):
        if name.endswith(suffix):
            return normalize_source_name(name[: -len(suffix)])
    return name


def domain(variable):
    if variable in {"Sex", "Age", "Marriage", "Course"}:
        return "Demographic and social"
    if variable == "BMI":
        return "Anthropometric"
    if variable in {"Oppinfection", "WHOstage", "HIVRNA", "CD4", "CD8"}:
        return "HIV disease"
    if variable in {"SCr", "Urea", "eGFR", "UACR"}:
        return "Renal"
    if variable in {"WBC", "PLT", "HB"}:
        return "Hematology"
    if variable in {"TC", "TG", "HDL", "LDL", "GLU", "HbA1c"}:
        return "Metabolic"
    if variable in {"ALT", "AST"}:
        return "Liver"
    if variable in {"Na", "P", "Ca", "K"}:
        return "Electrolyte and mineral"
    if variable == "CRP":
        return "Inflammation"
    if variable in {
        "CVD_status", "diabetes_status", "hypertension_status",
        "hypercholesterolemia_status", "HBV_status", "HCV_status",
    }:
        return "Comorbidity"
    if variable in {"antidiabetic_med", "antihypertensive_med", "antilipid_med"}:
        return "Concomitant medication"
    if variable.startswith("current_"):
        return "Current ART regimen"
    if variable.endswith("_cum_month"):
        return "Cumulative ART exposure"
    return "Other"


def read_csv_column(path, column):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return [row[column] for row in csv.DictReader(f)]


feature_groups = json.loads(FEATURE_GROUPS.read_text(encoding="utf-8"))
locked_final = set()
for key in ("continuous_vars", "binary_vars", "categorical_vars"):
    locked_final.update(normalize_source_name(x) for x in feature_groups[key])

rnn_source = {normalize_source_name(x) for x in read_csv_column(RNN_FEATURES, "feature_name")}
lstm_source = {normalize_source_name(x) for x in read_csv_column(LSTM_FEATURES, "feature_name")}

def dictionary_sources(path):
    with path.open("r", encoding="utf-8-sig", newline="") as f:
        return {normalize_source_name(row["source_variable"]) for row in csv.DictReader(f)}

cox_source = dictionary_sources(COX_FEATURES)
rsf_source = dictionary_sources(RSF_FEATURES)

wb = load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows = ws.iter_rows(values_only=True)
headers = [str(x) if x is not None else "" for x in next(rows)]
index = {name: i for i, name in enumerate(headers)}
candidate_vars = [name for name in headers if name and name not in NON_PREDICTORS]
dynamic_vars = [x for x in candidate_vars if x not in STATIC_VARS]

missing_rows = Counter()
observed_patients = defaultdict(set)
all_patients = set()
baseline_patients = set()
static_missing_patients = Counter()
row_n = 0

id_idx = index["ID"]
time_idx = index["time_bin"]
var_indices = {name: index[name] for name in candidate_vars}

for row in rows:
    row_n += 1
    raw_id = row[id_idx]
    if raw_id is None:
        continue
    patient_id = str(raw_id).strip()
    all_patients.add(patient_id)
    time_bin = row[time_idx]
    if time_bin == 0 or time_bin == 0.0 or str(time_bin).strip() == "0":
        baseline_patients.add(patient_id)
        for variable in STATIC_VARS:
            if is_missing(row[var_indices[variable]]):
                static_missing_patients[variable] += 1
    for variable in dynamic_vars:
        value = row[var_indices[variable]]
        if is_missing(value):
            missing_rows[variable] += 1
        else:
            observed_patients[variable].add(patient_id)

wb.close()

if baseline_patients != all_patients:
    raise RuntimeError(
        f"Baseline patient mismatch: all={len(all_patients)}, baseline={len(baseline_patients)}"
    )

records = []
for variable in candidate_vars:
    is_static = variable in STATIC_VARS
    if is_static:
        denominator = len(baseline_patients)
        missing_n = static_missing_patients[variable]
        never_pct = None
    else:
        denominator = row_n
        missing_n = missing_rows[variable]
        never_pct = 100.0 * (len(all_patients) - len(observed_patients[variable])) / len(all_patients)
    missing_pct = 100.0 * missing_n / denominator
    included = variable in locked_final
    if included:
        reason = ""
    elif variable in EXCLUDED_HIGH_MISSING:
        reason = "Substantial missingness / insufficient longitudinal availability"
    else:
        reason = "Manual confirmation required"
    records.append({
        "Variable": variable,
        "Clinical domain": domain(variable),
        "Variable type": "Baseline/static" if is_static else "Longitudinal/dynamic",
        "Missingness (%)": missing_pct,
        "Patients never observed (%)": never_pct,
        "Included in final model": "Yes" if included else "No",
        "Reason for exclusion": reason,
    })

records.sort(key=lambda x: (-x["Missingness (%)"], x["Variable"].lower()))

OUT_DIR.mkdir(parents=True, exist_ok=True)
with OUT_CSV.open("w", encoding="utf-8-sig", newline="") as f:
    writer = csv.DictWriter(f, fieldnames=list(records[0]))
    writer.writeheader()
    writer.writerows(records)

model_checks = {
    "locked_final_clinical_variables": sorted(locked_final),
    "locked_final_n": len(locked_final),
    "rnn_missing_from_locked": sorted(locked_final - rnn_source),
    "lstm_missing_from_locked": sorted(locked_final - lstm_source),
    "cox_missing_from_locked_after_reference_coding": sorted(locked_final - cox_source),
    "rsf_missing_from_locked_after_reference_coding": sorted(locked_final - rsf_source),
}

summary = {
    "source": str(SOURCE),
    "source_rows": row_n,
    "patients": len(all_patients),
    "baseline_patients": len(baseline_patients),
    "candidate_predictor_n": len(candidate_vars),
    "included_n": sum(r["Included in final model"] == "Yes" for r in records),
    "excluded_n": sum(r["Included in final model"] == "No" for r in records),
    "top10": records[:10],
    "model_checks": model_checks,
    "excluded": [r for r in records if r["Included in final model"] == "No"],
}
OUT_JSON.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")

print(json.dumps(summary, ensure_ascii=False, indent=2))
