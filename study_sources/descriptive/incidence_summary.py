from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


OUT = Path(r"__CKD_LOCAL_INPUT__/Documents/深度学习/outputs/ckd-incidence-cohort-comparison-20260905/Table_S_CKD_incidence_and_follow-up.docx")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_border(cell, *, top=None, bottom=None, left=None, right=None):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge, spec in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        if spec is None:
            continue
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), spec.get("val", "single"))
        node.set(qn("w:sz"), str(spec.get("sz", 6)))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), spec.get("color", "000000"))


def set_run_font(run, size=10, bold=False, italic=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def write_cell(cell, text, *, bold=False, align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(text)
    set_run_font(r, 10, bold=bold)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_margins(cell)


def main():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.27)
    sec.bottom_margin = Cm(1.27)
    sec.left_margin = Cm(1.27)
    sec.right_margin = Cm(1.27)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(10)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    run = title.add_run(
        "eTable X Incidence and follow-up of CKD in the development and external validation cohorts"
    )
    set_run_font(run, 10, bold=True)

    rows = [
        ("No. of participants", "31,911", "2,677"),
        ("Incident CKD events, n (%)", "2,491 (7.8)", "147 (5.5)"),
        ("Total follow-up, person-years", "193,725.0", "27,826.2"),
        ("Median follow-up, years [IQR]", "5.82 [2.89, 8.91]", "10.82 [9.33, 12.06]"),
        (
            "CKD incidence rate, per 1,000 person-years (95% CI)",
            "12.86 (12.36, 13.37)",
            "5.28 (4.46, 6.21)",
        ),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.style = "Normal Table"
    table.autofit = False
    widths = [Cm(7.2), Cm(5.9), Cm(5.9)]
    headers = [
        "Characteristic",
        "Shenzhen–Nanning development cohort",
        "Chongqing external validation cohort",
    ]
    for i, cell in enumerate(table.rows[0].cells):
        cell.width = widths[i]
        write_cell(cell, headers[i], bold=True)
        set_cell_shading(cell, "D9EAF7")
        set_cell_border(
            cell,
            top={"sz": 10, "color": "000000"},
            bottom={"sz": 6, "color": "000000"},
            left={"val": "nil"},
            right={"val": "nil"},
        )

    for ri, values in enumerate(rows):
        cells = table.add_row().cells
        for i, (cell, value) in enumerate(zip(cells, values)):
            cell.width = widths[i]
            write_cell(cell, value)
            set_cell_border(cell, left={"val": "nil"}, right={"val": "nil"})
        if ri == len(rows) - 1:
            for cell in cells:
                set_cell_border(cell, bottom={"sz": 10, "color": "000000"})

    note = doc.add_paragraph()
    note.paragraph_format.space_before = Pt(4)
    note.paragraph_format.space_after = Pt(0)
    note.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    note.paragraph_format.keep_together = True
    r = note.add_run("Notes: ")
    set_run_font(r, 10, bold=True)
    text = (
        "The development cohort comprised the pooled Shenzhen–Nanning population used for model "
        "development and internal validation; the independent external validation cohort comprised "
        "participants from Chongqing. Follow-up was calculated from ART initiation to the first incident "
        "CKD event or censoring at the last follow-up. Total person-time was calculated by summing each "
        "participant's follow-up interval in months and dividing by 12. Incidence rates are expressed per "
        "1,000 person-years with exact Poisson 95% confidence intervals. CKD, chronic kidney disease; "
        "ART, antiretroviral therapy; CI, confidence interval; IQR, interquartile range."
    )
    r = note.add_run(text)
    set_run_font(r, 10)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
